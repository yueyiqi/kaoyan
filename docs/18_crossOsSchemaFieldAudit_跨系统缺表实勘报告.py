"""Rebuild corrected field-report figures from evidence, never simulated labels.
[IoC_Skill_Telemetry] Status: PARTIAL; Author: Codex/OpenAI; Date: 2026-10-24.
"""
from pathlib import Path
import html
import json
import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET

BASE = Path(__file__).resolve().parent
ROOT = next(p for p in BASE.parents if (p / 'scripts/7_renderers_渲染中枢').exists())
sys.path.insert(0, str(ROOT / 'scripts/7_renderers_渲染中枢'))
from svg_bento_layout_engine import BentoGrid2x2, CardSpec
from pptx import Presentation
from pptx.util import Inches

NAME = Path(__file__).stem
OLD = '18_crossOsSnapshotArchive_快照阶段归档'
SWIM = '18_crossOsSnapshotSwimlane_快照阶段历史泳道'
EB = ROOT / '_ephemeral_blackbox/2026-10/eb_20261019_srbg10_H_codex_gpt-6_kaoyan-cross-os-schema-field-audit'
SVG_NS = 'http://www.w3.org/2000/svg'
ET.register_namespace('', SVG_NS)


def render_png(stem):
    renderer = shutil.which('rsvg-convert')
    if not renderer:
        raise RuntimeError('rsvg-convert is required to render the vector diagram')
    subprocess.run([renderer, '-o', str(BASE / (stem + '.png')), str(BASE / (stem + '.svg'))], check=True)


def overview():
    grid = BentoGrid2x2('跨系统缺表实勘 (Cross-OS Field Audit)',
                       '2026-10-24 · W1 / HP / S1 native runtimes; live failover remains unverified',
                       canvas_w=1800, canvas_h=1120)
    grid.add_cards([
        CardSpec('缺表根因', 'Missing schema root cause', '已修复 FIXED', accent='#2563eb', items=[
            ('迁移路径', 'Migration mount did not match the code'),
            ('真实故障', 'HP restart loop: schema_migrations missing'),
            ('修复结果', '14 tables ready; existing rows preserved')]),
        CardSpec('原生测试', 'Native runtime acceptance', '39 组 GROUPS', accent='#059669', items=[
            ('Windows', 'W1 native CPython: 13 tests'),
            ('Linux', 'HP native CPython: 13 tests'),
            ('macOS', 'S1 arm64: 13 tests; authorized replacement')]),
        CardSpec('数据往返', 'Full snapshot round trip', '一致 MATCH', accent='#7c3aed', items=[
            ('真实路径', 'W1 → HP → S1 → W1'),
            ('完整内容', 'Synthetic DB: 14 app + 1 extension table'),
            ('完整校验', 'Rows and full SHA-256 match on every hop')]),
        CardSpec('验收边界', 'Remaining acceptance', '部分 PARTIAL', accent='#d97706', items=[
            ('三机目标', 'W1 / HP / S1 snapshot tests accepted'),
            ('在线切换', 'No evidence for zero downtime or RPO=0'),
            ('PG18', 'PostgreSQL 18 replication is not verified')]),
    ])
    grid.render_to_file(BASE / (NAME + '.svg'))
    render_png(NAME)


def swimlane():
    svg = ET.Element('{' + SVG_NS + '}svg', {'viewBox': '0 0 1800 1120', 'width': '100%', 'preserveAspectRatio': 'xMidYMid meet'})
    def node(tag, attrs, text=None):
        item = ET.SubElement(svg, tag, {k: str(v) for k, v in attrs.items()})
        item.text = text
        return item
    def text(x, y, value, size=24, color='#0f172a', anchor='start'):
        node('text', {'x': x, 'y': y, 'font-size': size, 'fill': color, 'text-anchor': anchor,
                      'font-family': 'PingFang SC, Microsoft YaHei, sans-serif'}, value)
    node('rect', {'width': 1800, 'height': 1120, 'fill': '#f8fafc'})
    text(60, 65, '跨系统快照迁移实测泳道 (Cross-OS Snapshot Migration)', 36)
    text(60, 108, 'W1 → HP → S1 → W1 · 实机逐表校验 / Actual hosts and full table checksums', 23, '#475569')
    text(60, 149, '用户授权 S1 替代 M1；快照迁移已实测，在线切主仍待验收', 24, '#b45309')
    columns = [240, 660, 1080, 1500]
    for x, title, detail, color in zip(columns,
            ['W1 · Windows', 'HP · Linux', 'S1 · macOS', '校验 · Verification'],
            ['AMD64 / Native Python', 'x86_64 / Native Python', 'arm64 / Native Python', '15 表 / Full inventory'],
            ['#2563eb', '#059669', '#7c3aed', '#b45309']):
        text(x, 220, title, 29, color, 'middle')
        text(x, 257, detail, 20, '#475569', 'middle')
        node('line', {'x1': x, 'x2': x, 'y1': 280, 'y2': 880, 'stroke': '#cbd5e1', 'stroke-width': 2, 'stroke-dasharray': '8 8'})
    steps = [(0, 1, 350, '#2563eb', '1. 完整备份包 (Full snapshot)', '合成库：14 应用表 + 1 扩展表；保留 WAL 提交'),
             (1, 2, 500, '#059669', '2. 恢复与再导出 (Restore and export)', '不同 SQLite 版本：文件校验 + 完整性 + 逐表摘要'),
             (2, 0, 650, '#7c3aed', '3. 返回 Windows (Return to Windows)', '中文、音标、CRLF、路径文本和 BLOB 均保持原值')]
    for start, end, y, color, title, detail in steps:
        x1, x2 = columns[start], columns[end]
        node('line', {'x1': x1, 'x2': x2, 'y1': y, 'y2': y, 'stroke': color, 'stroke-width': 4})
        sign = 1 if x2 > x1 else -1
        node('polygon', {'points': f'{x2},{y} {x2-sign*18},{y-9} {x2-sign*18},{y+9}', 'fill': color})
        center = (x1 + x2) / 2
        text(center, y-48, title, 25, color, 'middle')
        text(center, y-17, detail, 20, '#334155', 'middle')
    text(columns[0], 790, '4. 全表摘要一致', 26, '#059669', 'middle')
    text(columns[0], 827, 'All table hashes match', 20, '#475569', 'middle')
    text(columns[3], 440, '逐跳校验均通过', 26, '#b45309', 'middle')
    text(columns[3], 481, 'SHA-256 / Rows / Columns', 20, '#475569', 'middle')
    node('rect', {'x': 60, 'y': 914, 'width': 1680, 'height': 146, 'rx': 18, 'fill': '#fff7ed', 'stroke': '#fed7aa'})
    text(90, 957, '切换契约 (Cutover): 停止源写入 → 导出 → 验证 → 新库恢复 → 业务验收 → 切换入口', 24, '#9a3412')
    text(90, 1001, '已验证快照可移植性；在线复制、PG18、零停机与 RPO=0 均未验收。', 24, '#9a3412')
    text(60, 1096, '证据 Evidence: ev-0007 / ev-0039 / ev-0040 / ev-0041 / ev-0024 · srbg10 / Codex · 2026-10-24', 18, '#64748b')
    ET.ElementTree(svg).write(BASE / (SWIM + '.svg'), encoding='utf-8', xml_declaration=True)
    render_png(SWIM)


def derivatives():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(15), Inches(9.333)
    for stem in [NAME, SWIM]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(BASE / (stem + '.png')), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(BASE / (NAME + '.pptx'))
    for suffix in ['svg', 'png', 'pptx']:
        shutil.copy2(BASE / (NAME + '.' + suffix), BASE / (OLD + '.' + suffix))
    for suffix in ['svg', 'png']:
        shutil.copy2(BASE / (NAME + '.' + suffix), BASE / (OLD + '_slides_preview.' + suffix))
    (BASE / (OLD + '_slides_preview.md')).write_text(
        '# 修订后的幻灯片预览 / Corrected slide preview\n\nStatus: PARTIAL\n\n'
        '两页：实地结果概览；真实快照往返泳道。W1/HP/S1 快照已验收；PG18 和在线切主未验收。\n\n'
        'Evidence: ev:ev-0024; ev:ev-0043.\n', encoding='utf-8')
    body = '<h1>跨系统缺表实勘 (Cross-OS Field Audit)</h1><p>PARTIAL · W1/HP/S1 快照已验收；PG18 与在线切主未验收。</p>'
    for stem in [NAME, SWIM]:
        body += '<img alt="实测结果 (Measured evidence)" src="' + stem + '.svg">'
    page = '''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<script>document.documentElement.dataset.theme=localStorage.getItem('field-theme')||'daylight'</script>
<style>:root{--bg:#f8fafc;--fg:#172033}[data-theme=midnight]{--bg:#111827;--fg:#f8fafc}[data-theme=eyecare]{--bg:#edf3e5;--fg:#26352a}body{background:var(--bg);color:var(--fg);font:17px system-ui;max-width:1500px;margin:auto;padding:24px}img{width:100%;height:auto}button{padding:12px;margin:4px}@media(max-width:700px){body{padding:12px;font-size:15px}}</style>
<button onclick="setTheme('daylight')">日光 Daylight</button><button onclick="setTheme('midnight')">夜间 Midnight</button><button onclick="setTheme('eyecare')">护眼 Eyecare</button>'''+body+'''
<script>function setTheme(t){document.documentElement.dataset.theme=t;localStorage.setItem('field-theme',t)}</script></html>'''
    for stem in [NAME, OLD]:
        (BASE / (stem + '.html')).write_text(page, encoding='utf-8')


if __name__ == '__main__':
    evidence = [json.loads(line) for line in (EB / 'evidence.jsonl').read_text().splitlines()]
    for ident in ['ev-0007', 'ev-0039', 'ev-0040', 'ev-0041', 'ev-0024']:
        assert any(r['id'] == ident and r['exit'] == 0 for r in evidence), ident
    overview()
    swimlane()
    derivatives()
    print('Corrected report figures generated; W1/HP/S1 snapshot verified; online failover remains PARTIAL')

# [IoC_Revision_v2] Status: VERIFIED for W1/HP/S1 snapshots, ev:ev-0024.
