"""Render measured live-switchover report and replace the original paper-only swimlane.
[IoC_Skill_Telemetry] Status: VERIFIED; ev:ev-0029; srbg10 / Codex.
"""
from pathlib import Path
import json
import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET
BASE=Path(__file__).resolve().parent
ROOT=next(p for p in BASE.parents if (p/'scripts/7_renderers_渲染中枢').exists())
EB=ROOT/'_ephemeral_blackbox/2026-10/eb_20261020_srbg10_H_codex_gpt-6_kaoyan-live-switchover-windows'
sys.path.insert(0,str(ROOT/'scripts/7_renderers_渲染中枢'))
from svg_bento_layout_engine import BentoGrid2x2,CardSpec
from pptx import Presentation
from pptx.util import Inches
NAME=Path(__file__).stem
SWIM='17_crossOsHaRoamingSwimlane_跨系统瞬时迁移泳道图'
PRD='19_liveSwitchoverPrd_在线切主语义单与实施方案'
def png(stem):
    subprocess.run([shutil.which('rsvg-convert'),'-o',str(BASE/(stem+'.png')),str(BASE/(stem+'.svg'))],check=True)
def overview():
    grid=BentoGrid2x2('在线切主实机验收 (Live Switchover Acceptance)',
        'W1 Windows · HP Linux · S1 macOS · production evidence, 2026-10-24',canvas_w=1800,canvas_h=1120)
    grid.add_cards([
        CardSpec('Windows 原生落库','Native Windows database','已验收 VERIFIED',accent='#2563eb',items=[
            ('常驻服务','SYSTEM startup task; native CPython'),
            ('生产库表','18 tables: 14 app + 2 legacy + 2 protocol'),
            ('真实恢复','Dropped populated table; full data restored')]),
        CardSpec('公网真实切主','Public production switchover','已验收 VERIFIED',accent='#059669',items=[
            ('执行路线','HP → W1 → S1 → HP'),
            ('一致性','Historical rows and full DDL preserved'),
            ('原始证据','Public HTTPS writes and replay receipts')]),
        CardSpec('故障与修复','Faults and repairs','已复测 RETESTED',accent='#7c3aed',items=[
            ('隔离保护','Epoch fencing; two durable replica ACKs'),
            ('故障注入','Link loss, worker exit, missing Windows table'),
            ('修复闭环','Schema, ownership, bad input, status, proxy')]),
        CardSpec('实测边界','Measured boundaries','如实披露 MEASURED',accent='#b45309',items=[
            ('首次迁入','24.79 seconds of initial write freeze'),
            ('切主耗时','Seconds, not the old 18ms paper claim'),
            ('可用性','One coordinator; no quorum means no ACK')])])
    grid.render_to_file(BASE/(NAME+'.svg'));png(NAME)
def swimlane():
    ns='http://www.w3.org/2000/svg'
    ET.register_namespace('',ns)
    root=ET.Element('{'+ns+'}svg',{'viewBox':'0 0 1800 1240','width':'100%','preserveAspectRatio':'xMidYMid meet'})
    def element(tag,attrs,text=None):
        item=ET.SubElement(root,tag,{k:str(v) for k,v in attrs.items()});item.text=text
    def text(x,y,value,size=23,color='#0f172a',anchor='start'):
        element('text',{'x':x,'y':y,'font-size':size,'fill':color,'text-anchor':anchor,
            'font-family':'PingFang SC, Microsoft YaHei, sans-serif'},value)
    element('rect',{'width':1800,'height':1240,'fill':'#f8fafc'})
    text(60,65,'跨系统在线切主：已部署与实测 (Cross-OS Live Switchover)',35)
    text(60,108,'生产库 18 表；S1 经用户授权替代 M1 · 18 production tables; native Windows / Linux / macOS',22,'#475569')
    text(60,150,'示例：W1 执行主，HP 协调入口；生产最终主节点 HP (Example: W1 primary; final primary HP)',22,'#475569')
    xs=[210,660,1110,1560]
    titles=['客户端 / Client','HP / Linux','W1 / Windows','S1 / macOS']
    subtitles=['同一 HTTPS 入口','协调器 + 本地副本','原生 SQL 执行主示例','原生完整副本']
    for x,title,sub in zip(xs,titles,subtitles):
        text(x,224,title,27,'#1d4ed8','middle');text(x,260,sub,21,'#475569','middle')
        element('line',{'x1':x,'x2':x,'y1':280,'y2':970,'stroke':'#cbd5e1','stroke-width':2,'stroke-dasharray':'7 7'})
    steps=[
        (0,1,350,'1. POST + 同一请求键 (Stable request key)','#2563eb'),
        (1,2,470,'2. epoch 校验 + 整请求事务 (Fence + atomic transaction)','#059669'),
        (2,1,590,'3. 完整快照 + 回执 (Snapshot + response receipt)','#7c3aed'),
        (1,3,710,'4. 内容与结构校验；至少两机确认 (Verify + 2 ACKs)','#b45309'),
        (1,0,830,'5. 持久化后返回成功 (Reply only after durable ACKs)','#2563eb')]
    for a,b,y,label,color in steps:
        x1,x2=xs[a],xs[b];direction=1 if x2>x1 else -1
        text(390,y-35,label,23,color)
        element('line',{'x1':x1,'x2':x2,'y1':y,'y2':y,'stroke':color,'stroke-width':4})
        element('polygon',{'points':f'{x2},{y} {x2-direction*18},{y-9} {x2-direction*18},{y+9}','fill':color})
    text(60,938,'切主：排队 → 持久化新主及任期 → 校验副本 → 重放未确认请求 → 继续同入口写入',24,'#0f172a')
    element('rect',{'x':60,'y':985,'width':1680,'height':170,'rx':18,'fill':'#fff7ed','stroke':'#fed7aa'})
    text(88,1029,'实际通过：公网 HP→W1→S1→HP；跨机隧道中断；W1 缺表恢复；全部历史行保留',25,'#9a3412')
    text(88,1071,'边界：首次迁入冻结 24.79 秒；后续切主为秒级；少于两个副本时拒绝确认；HP 协调器仍是单点',23,'#9a3412')
    text(88,1113,'Verified native paths; initial freeze 24.79s; seconds per switch; single coordinator remains a dependency',20,'#9a3412')
    text(60,1202,'Evidence: ev-0021 / ev-0023 / ev-0029 · srbg10 / Codex · 2026-10-24',19,'#64748b')
    ET.ElementTree(root).write(BASE/(SWIM+'.svg'),encoding='utf-8',xml_declaration=True);png(SWIM)
def main():
    ledger=[json.loads(line) for line in (EB/'evidence.jsonl').read_text().splitlines()]
    for ident in ['ev-0021','ev-0023','ev-0029']:
        assert any(r['id']==ident and r['exit']==0 for r in ledger),ident
    overview();swimlane()
    prs=Presentation();prs.slide_width=Inches(15);prs.slide_height=Inches(10.333)
    for stem in [NAME,SWIM]:
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(BASE/(stem+'.png')),0,0,width=prs.slide_width)
    prs.save(BASE/(NAME+'.pptx'))
    for target in [PRD,'17_kaoyanCrossOsRoamingHaPrd_跨系统漫游高可用与验收','README']:
        for suffix in ['svg','png','pptx']:
            shutil.copy2(BASE/(NAME+'.'+suffix),BASE/(target+'.'+suffix))
    print('Measured acceptance SVG/PNG/PPTX and original swimlane regenerated')
if __name__=='__main__':main()

